import unittest
import shutil
from pathlib import Path
import docx

from backend.app.services.checkpoint import (
    CheckpointVendorPMParser,
    CheckpointSectionCropper,
    CheckpointPMEvaluator,
    CheckpointCTCReportGenerator,
    CheckPointReportData,
)

class TestCheckpointVendorPMConversion(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.vendor_pm_dir = Path("doc/CIMB_3rd_Quater_of_2026_PM/Vendor PM")
        cls.target_docx = Path("doc/CIMB_3rd_Quater_of_2026_PM/CIMB_3rd_Quater_of_2026_Preventive_Maintenance_Report.docx")
        cls.workspace_test_dir = Path("workspace/tests/checkpoint_pm")
        cls.workspace_test_dir.mkdir(parents=True, exist_ok=True)

        cls.parser = CheckpointVendorPMParser()
        cls.cropper = CheckpointSectionCropper(dpi=150)
        cls.evaluator = CheckpointPMEvaluator()
        cls.generator = CheckpointCTCReportGenerator()

        # Parse all 4 devices
        cls.pdf_files = sorted(cls.vendor_pm_dir.glob("*.pdf"))
        cls.devices = []
        for p in cls.pdf_files:
            dev = cls.parser.parse_pdf(p)
            cls.evaluator.evaluate(dev)
            dev_evidence_dir = cls.workspace_test_dir / "evidence" / dev.system.hostname
            cls.cropper.crop_device_evidence(p, dev, dev_evidence_dir)
            cls.devices.append(dev)

    def test_parsed_devices_count(self):
        self.assertEqual(len(self.devices), 4)

    def test_telemetry_accuracy(self):
        dev_map = {d.system.hostname: d for d in self.devices}

        # 1. Ki-L10-B2-FW01
        fw01 = dev_map.get("Ki-L10-B2-FW01")
        self.assertIsNotNone(fw01)
        self.assertEqual(fw01.system.model, "Check Point 6900")
        self.assertEqual(fw01.system.serial, "2201BA1755")
        self.assertEqual(fw01.system.mgmt_ip, "10.42.2.253")
        self.assertTrue(fw01.system.is_gateway)
        self.assertFalse(fw01.storage_log.var_log_full)
        self.assertTrue(fw01.storage_log.log_switch_enabled)

        # 2. Ki-L10-B3-FW02
        fw02 = dev_map.get("Ki-L10-B3-FW02")
        self.assertIsNotNone(fw02)
        self.assertEqual(fw02.system.model, "Check Point 6900")
        self.assertEqual(fw02.system.serial, "2201BA1778")
        self.assertEqual(fw02.system.mgmt_ip, "10.42.2.254")
        self.assertTrue(fw02.system.is_gateway)
        self.assertFalse(fw02.storage_log.var_log_full)
        self.assertTrue(fw02.storage_log.log_switch_enabled)

        # 3. Ki-L10-B2-MGMT-FW01
        mgmt01 = dev_map.get("Ki-L10-B2-MGMT-FW01")
        self.assertIsNotNone(mgmt01)
        self.assertIn("Smart-1", mgmt01.system.model)
        self.assertEqual(mgmt01.system.serial, "51WX6L3")
        self.assertEqual(mgmt01.system.mgmt_ip, "10.42.2.251")
        self.assertFalse(mgmt01.system.is_gateway)
        self.assertTrue(mgmt01.storage_log.var_log_full)
        self.assertFalse(mgmt01.storage_log.log_switch_enabled)

        # 4. Ki-L10-B3-MGMT-FW02
        mgmt02 = dev_map.get("Ki-L10-B3-MGMT-FW02")
        self.assertIsNotNone(mgmt02)
        self.assertIn("Smart-1", mgmt02.system.model)
        self.assertEqual(mgmt02.system.serial, "41V92L3")
        self.assertEqual(mgmt02.system.mgmt_ip, "10.42.2.252")
        self.assertFalse(mgmt02.system.is_gateway)
        self.assertFalse(mgmt02.storage_log.var_log_full)
        self.assertFalse(mgmt02.storage_log.log_switch_enabled)

    def test_checklist_evaluations(self):
        dev_map = {d.system.hostname: d for d in self.devices}

        # Gateways: Item 11 -> App 12, Item 12 -> App 13
        fw01 = dev_map["Ki-L10-B2-FW01"]
        cl_fw01 = {item.no: item for item in fw01.checklist}
        self.assertEqual(cl_fw01[1].pass_status, "Y")
        self.assertEqual(cl_fw01[2].pass_status, "N/A")
        self.assertEqual(cl_fw01[8].pass_status, "Y")
        self.assertEqual(cl_fw01[9].pass_status, "Y")
        self.assertIn("Enabled", cl_fw01[9].remarks)
        self.assertIn("Appendix 12", cl_fw01[11].remarks)
        self.assertIn("Appendix 13", cl_fw01[12].remarks)

        # Management 1: Item 8 -> Pass=N (/var/log 100% full), Item 9 -> Disabled
        mgmt01 = dev_map["Ki-L10-B2-MGMT-FW01"]
        cl_mgmt01 = {item.no: item for item in mgmt01.checklist}
        self.assertEqual(cl_mgmt01[8].pass_status, "N")
        self.assertIn("/var/log 100% full", cl_mgmt01[8].remarks)
        self.assertIn("Disabled", cl_mgmt01[9].remarks)
        self.assertIn("Appendix 11", cl_mgmt01[11].remarks)
        self.assertIn("Appendix 12", cl_mgmt01[12].remarks)
        self.assertIn("/var/log Directory Full", mgmt01.key_highlights)

    def test_appendix_evidence_crops(self):
        dev_map = {d.system.hostname: d for d in self.devices}

        for host, dev in dev_map.items():
            self.assertGreater(len(dev.appendix_images), 10)
            self.assertIsNotNone(dev.recommendation_crop_path)
            self.assertTrue(Path(dev.recommendation_crop_path).exists())

            # Check all image paths exist on disk
            for img in dev.appendix_images:
                p = Path(img.image_path)
                self.assertTrue(p.exists(), f"Image not found: {p}")
                self.assertGreater(p.stat().st_size, 0)

            # Gateways should cover appendix 1..13
            if dev.system.is_gateway:
                apps = {img.appendix_no for img in dev.appendix_images}
                self.assertTrue({1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13}.issubset(apps))
            else:
                apps = {img.appendix_no for img in dev.appendix_images}
                self.assertTrue({1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12}.issubset(apps))

    def test_full_report_generation(self):
        out_doc_path = self.workspace_test_dir / "Generated_CIMB_Report.docx"
        report_data = CheckPointReportData(
            customer="CIMB BANK BERHAD",
            report_title="3rd Quarter Of 2026 Preventive Maintenance Report For Check Point Firewall",
            author_name="Cha Zi Yu",
            author_designation="Associate Technical Engineer",
            devices=self.devices,
        )

        final_path = self.generator.generate_report(
            report_data=report_data,
            output_path=out_doc_path,
            template_path=self.target_docx,
        )

        self.assertTrue(final_path.exists())
        self.assertGreater(final_path.stat().st_size, 5_000_000)

        # Inspect generated docx tables
        doc = docx.Document(str(final_path))
        self.assertEqual(len(doc.tables), 23)

        # Verify Device 1 (FW01)
        t_app1 = doc.tables[2]
        self.assertEqual(t_app1.rows[1].cells[1].text, "2201BA1755")
        self.assertEqual(t_app1.rows[2].cells[1].text, "10.42.2.253")

        # Verify Device 3 (MGMT01 - var/log full)
        t_app3 = doc.tables[13]
        self.assertEqual(t_app3.rows[1].cells[1].text, "51WX6L3")
        self.assertEqual(t_app3.rows[2].cells[1].text, "10.42.2.251")
        self.assertIn("/var/log Directory Full", doc.tables[14].rows[0].cells[0].text)

if __name__ == "__main__":
    unittest.main()
