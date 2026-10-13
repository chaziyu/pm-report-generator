"""
Comprehensive Unit & Integration Test Suite for Trend Micro Deep Security PM Pipeline.
Tests dataset importer, expected value extractor, rules evaluator, report generator,
SFT/few-shot training dataset pairs, and Local AI prompt integration.
"""

import json
import unittest
from pathlib import Path
import docx
from docx.oxml.ns import qn

from backend.app.services.deep_security import (
    DeepSecurityReportData,
    DeepSecuritySystemData,
    DeepSecurityServerHealth,
    DeepSecurityLicenseData,
    DeepSecurityAlertsEvents,
    DeepSecurityEvaluator,
    DeepSecurityCTCReportGenerator,
    DeepSecurityReportImporter,
    DeepSecurityParser,
    DEEP_SECURITY_APPENDIX_MAP,
)
from backend.app.services.ai.prompts import (
    DEEP_SECURITY_ALLOWED_CATEGORIES,
    DEEP_SECURITY_PM_SYSTEM_PROMPT,
    build_deep_security_pm_prompt,
)

class TestDeepSecurityPMPipeline(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.source_docx = Path("doc/LQ Retail_2026_Preventive_Maintenance.docx")
        cls.test_workspace = Path("workspace/tests/deep_security_pm")
        cls.test_workspace.mkdir(parents=True, exist_ok=True)
        cls.dataset_dir = cls.test_workspace / "dataset"
        cls.importer = DeepSecurityReportImporter(output_dir=cls.dataset_dir, overwrite=True)

    def test_importer_and_dataset_structure(self):
        """Verifies report ingestion, image extraction, and dataset structure."""
        self.assertTrue(self.source_docx.exists())
        res = self.importer.import_report(self.source_docx)

        case_id = res["case_id"]
        case_dir = self.dataset_dir / case_id
        self.assertTrue(case_dir.exists())

        # Check expected.json and manifest.json
        expected_path = case_dir / "expected.json"
        manifest_path = case_dir / "manifest.json"
        self.assertTrue(expected_path.exists())
        self.assertTrue(manifest_path.exists())

        with open(manifest_path, "r", encoding="utf-8") as f:
            manifest = json.load(f)
        self.assertEqual(manifest["template_variant"], "trend_micro_deep_security_pm_v1")
        self.assertEqual(len(manifest["appendices"]), 9)

        # Check extracted appendix screenshots
        app_dir = case_dir / "appendices"
        extracted_imgs = list(app_dir.glob("*.png"))
        self.assertGreaterEqual(len(extracted_imgs), 14)
        for img in extracted_imgs:
            self.assertGreater(img.stat().st_size, 1000)

    def test_expected_values_extraction(self):
        """Verifies ground truth values extracted from LQ Retail report."""
        doc = docx.Document(str(self.source_docx))
        data = self.importer.extract_expected_values(doc)

        meta = data["metadata"]
        app = data["application"]
        self.assertIn("LQ Retail", meta["customer"])
        self.assertEqual(meta["fsms_id"], "02892342")
        self.assertEqual(meta["contract_no"], "ISS/2025/M4170")
        self.assertEqual(meta["contract_period"], "01/04/2026 - 31/03/2027")
        self.assertEqual(meta["pm_date"], "20/05/2026")
        self.assertEqual(meta["author_name"], "Loh Zi Xin")
        self.assertEqual(meta["verified_name"], "Fatin Amirah Binti Che A Latif")

        self.assertEqual(app["model"], "Deep Security")
        self.assertEqual(app["hostname"], "DSM-GF-MSER-01")
        self.assertEqual(app["manager_version"], "20.0.844")

        # Verify checklist extraction
        checklist = data["checklist"]
        self.assertGreaterEqual(len(checklist), 12)
        # All items should have passed
        for item in checklist:
            self.assertEqual(item["result"], "Yes")

    def test_evaluator_and_rules(self):
        """Verifies evaluation rules, health thresholds, and upgrade recommendations."""
        parser = DeepSecurityParser()
        evaluator = DeepSecurityEvaluator()

        report_data = parser.parse_telemetry({
            "system": {
                "hostname": "DSM-GF-MSER-01",
                "manager_version": "20.0.844",
                "recommended_manager_version": "20.0.1160",
                "agent_version": "20.0.0.8137",
                "recommended_agent_version": "20.0.3.8610",
            },
            "health": {
                "disk_space_ok": True,
                "server_performance_healthy": True,
            },
        })

        evaluator.evaluate(report_data)

        # 4 Highlights
        self.assertEqual(len(report_data.key_highlights), 4)
        self.assertIn("operating normally", report_data.key_highlights[0])

        # 2 Recommendations
        self.assertEqual(len(report_data.recommendations), 2)
        rec_mgr = report_data.recommendations[0]
        self.assertEqual(rec_mgr.item, "Deep Security Manager Version")
        self.assertEqual(rec_mgr.current_version, "20.0.844")
        self.assertEqual(rec_mgr.recommended_version, "20.0.1160")

        # 9 Checklist items
        self.assertEqual(len(report_data.checklist), 9)
        self.assertEqual(report_data.checklist[0].task, "Dashboard and Widgets")
        self.assertEqual(report_data.checklist[8].task, "Deep Security Manager and Database Server Health Check")
        self.assertEqual(report_data.checklist[8].subtasks[0].result, "Yes")

    def test_full_report_generation(self):
        """Verifies that generated Word DOCX opens, has valid XML, and correct FORMCHECKBOX states."""
        parser = DeepSecurityParser()
        evaluator = DeepSecurityEvaluator()
        generator = DeepSecurityCTCReportGenerator()

        report_data = parser.parse_telemetry({})
        evaluator.evaluate(report_data)

        out_docx = self.test_workspace / "Generated_Deep_Security_Report.docx"
        res_path = generator.generate_report(report_data, out_docx)

        self.assertTrue(res_path.exists())
        self.assertGreater(res_path.stat().st_size, 3_000_000)

        # Inspect generated docx
        doc = docx.Document(str(res_path))
        self.assertEqual(len(doc.tables), 8)
        self.assertEqual(len(doc.paragraphs), 74)

        # Check Table 2 Model and Hostname
        t2 = doc.tables[2]
        self.assertEqual(t2.rows[0].cells[1].text.strip(), "Deep Security")
        self.assertEqual(t2.rows[1].cells[1].text.strip(), "DSM-GF-MSER-01")

        # Check Table 5 Checklist Checkbox
        t5 = doc.tables[5]
        subtask_row = t5.rows[2]  # Task 1.1
        cbs = subtask_row.cells[4]._tc.xpath(".//w:checkBox")
        self.assertEqual(len(cbs), 2)
        def0 = cbs[0].find(qn("w:default")).attrib.get(qn("w:val"))
        def1 = cbs[1].find(qn("w:default")).attrib.get(qn("w:val"))
        self.assertEqual(def0, "true")
        self.assertEqual(def1, "false")

    def test_training_pairs_jsonl_export(self):
        """Verifies training dataset pairs export in JSONL format."""
        jsonl_path = self.test_workspace / "training_pairs.jsonl"
        case_id = "case_02892342_deep_security"
        self.importer.export_training_dataset(case_id, jsonl_path)

        self.assertTrue(jsonl_path.exists())
        with open(jsonl_path, "r", encoding="utf-8") as f:
            lines = [line.strip() for line in f if line.strip()]

        self.assertGreaterEqual(len(lines), 1)
        record = json.loads(lines[0])
        self.assertIn("prompt", record)
        self.assertIn("completion", record)

        target = json.loads(record["completion"])
        self.assertEqual(target["model"], "Deep Security")
        self.assertEqual(target["hostname"], "DSM-GF-MSER-01")
        self.assertEqual(target["manager_version"], "20.0.844")
        self.assertEqual(target["appendices_count"], 9)
        self.assertGreaterEqual(len(target["checklist"]), 12)

    def test_local_ai_prompts(self):
        """Verifies Local AI prompt templates and allowed categories."""
        self.assertIn("dashboard_widgets", DEEP_SECURITY_ALLOWED_CATEGORIES)
        self.assertIn("server_health", DEEP_SECURITY_ALLOWED_CATEGORIES)
        self.assertIn("DSM-GF-MSER-01", DEEP_SECURITY_PM_SYSTEM_PROMPT)

        prompt = build_deep_security_pm_prompt("Sample OCR evidence text from Deep Security Manager")
        self.assertIn("DEEP SECURITY EVIDENCE TEXT", prompt)
        self.assertIn("Sample OCR evidence", prompt)

if __name__ == "__main__":
    unittest.main()
