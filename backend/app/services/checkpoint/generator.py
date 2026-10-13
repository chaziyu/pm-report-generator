import shutil
from pathlib import Path
from typing import List, Optional
import docx
from docx.shared import Inches, Pt, RGBColor

from .models import (
    CheckPointReportData,
    CheckPointDeviceData,
)

# Standard table indices in CIMB 4-device template
DEVICE_TABLE_MAPS = [
    {
        "device_name": "Ki-L10-B2-FW01",
        "appliance_table_idx": 2,
        "highlights_table_idx": 3,
        "checklist_table_indices": [4, 5],
        "appendix_table_indices": [6],
    },
    {
        "device_name": "Ki-L10-B3-FW02",
        "appliance_table_idx": 7,
        "highlights_table_idx": 8,
        "checklist_table_indices": [9, 10],
        "appendix_table_indices": [11, 12],
    },
    {
        "device_name": "Ki-L10-B2-MGMT-FW01",
        "appliance_table_idx": 13,
        "highlights_table_idx": 14,
        "checklist_table_indices": [15],
        "appendix_table_indices": [16, 17, 18],
    },
    {
        "device_name": "Ki-L10-B3-MGMT-FW02",
        "appliance_table_idx": 19,
        "highlights_table_idx": 20,
        "checklist_table_indices": [21],
        "appendix_table_indices": [22],
    },
]

class CheckpointCTCReportGenerator:
    """
    Generates official CTC Format Check Point Preventive Maintenance Word (.docx) reports.
    Preserves all corporate styles, headers, footers, tables, and embeds rendered evidence.
    """

    def generate_report(
        self,
        report_data: CheckPointReportData,
        output_path: Path,
        template_path: Optional[Path] = None,
    ) -> Path:
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)

        if template_path is None:
            template_path = Path("doc/CIMB_3rd_Quater_of_2026_PM/CIMB_3rd_Quater_of_2026_Preventive_Maintenance_Report.docx")

        if not template_path.exists():
            raise FileNotFoundError(f"Base CTC template not found at {template_path}")

        # Fresh copy from base template
        shutil.copyfile(str(template_path), str(output_path))
        doc = docx.Document(str(output_path))

        # 1. Update cover & metadata
        self._update_metadata(doc, report_data)

        # 2. Update each device section
        # Match by hostname first if available, otherwise by gateway/management topology
        dev_by_name = {dev.system.hostname.lower(): dev for dev in report_data.devices}
        assigned_devices = set()

        for slot_idx, table_map in enumerate(DEVICE_TABLE_MAPS):
            target_name = table_map["device_name"].lower()
            matched_dev = None

            if target_name in dev_by_name and target_name not in assigned_devices:
                matched_dev = dev_by_name[target_name]
                assigned_devices.add(target_name)
            else:
                # Find an unassigned device of matching role
                is_mgmt_slot = "mgmt" in target_name
                for dev in report_data.devices:
                    dev_key = dev.system.hostname.lower()
                    if dev_key not in assigned_devices:
                        dev_is_mgmt = not dev.system.is_gateway
                        if dev_is_mgmt == is_mgmt_slot:
                            matched_dev = dev
                            assigned_devices.add(dev_key)
                            break

            if not matched_dev and slot_idx < len(report_data.devices):
                matched_dev = report_data.devices[slot_idx]

            if matched_dev:
                self._populate_device_in_template(doc, matched_dev, table_map)

        doc.save(str(output_path))
        return output_path

    def _update_metadata(self, doc: docx.Document, report_data: CheckPointReportData) -> None:
        # Table 0 is Author Details
        if len(doc.tables) > 0:
            t0 = doc.tables[0]
            # Search for author name row
            for r in t0.rows:
                row_txt = " ".join(c.text for c in r.cells)
                if "Name" in row_txt and len(r.cells) >= 3:
                    r.cells[2].text = report_data.author_name
                elif "Designation" in row_txt and len(r.cells) >= 3:
                    r.cells[2].text = report_data.author_designation

    def _populate_device_in_template(
        self,
        doc: docx.Document,
        device: CheckPointDeviceData,
        table_map: dict
    ) -> None:
        sys = device.system

        # 1. Appliance Info Table
        app_tbl_idx = table_map["appliance_table_idx"]
        if app_tbl_idx < len(doc.tables):
            t_app = doc.tables[app_tbl_idx]
            for r in t_app.rows:
                lbl = r.cells[0].text.strip().lower()
                if "model" in lbl and len(r.cells) >= 2:
                    r.cells[1].text = sys.model or "Check Point"
                elif "serial" in lbl and len(r.cells) >= 2:
                    r.cells[1].text = sys.serial or ""
                elif "ip address" in lbl and len(r.cells) >= 2:
                    r.cells[1].text = sys.mgmt_ip or ""
                elif "date pm" in lbl and len(r.cells) >= 2:
                    r.cells[1].text = sys.date_pm or "23rd September 2026"

        # 2. Highlights & Recommendation Table
        hl_tbl_idx = table_map["highlights_table_idx"]
        if hl_tbl_idx < len(doc.tables):
            t_hl = doc.tables[hl_tbl_idx]
            cell = t_hl.rows[0].cells[0]
            self._populate_highlights_cell(cell, device)

        # 3. Checklist Tables
        cl_indices = table_map["checklist_table_indices"]
        cl_tables = [doc.tables[idx] for idx in cl_indices if idx < len(doc.tables)]
        self._populate_checklist_tables(cl_tables, device)

        # 4. Appendix Tables
        app_indices = table_map["appendix_table_indices"]
        app_tables = [doc.tables[idx] for idx in app_indices if idx < len(doc.tables)]
        self._populate_appendix_tables(app_tables, device)

    def _populate_highlights_cell(self, cell, device: CheckPointDeviceData) -> None:
        # Clear paragraphs
        p0 = cell.paragraphs[0]
        p0.text = ""
        for p in list(cell.paragraphs[1:]):
            p._p.getparent().remove(p._p)

        # Add Key Highlights
        p_hl_title = cell.paragraphs[0]
        r1 = p_hl_title.add_run("Key Highlights\n")
        r1.bold = True
        r1.font.name = "Calibri"
        r1.font.size = Pt(11)

        for hl in device.key_highlights:
            p_hl = cell.add_paragraph()
            r = p_hl.add_run(hl)
            r.font.name = "Calibri"
            r.font.size = Pt(10)

        # Add Recommendation
        p_rec_title = cell.add_paragraph()
        r2 = p_rec_title.add_run("\nRecommendation")
        r2.bold = True
        r2.font.name = "Calibri"
        r2.font.size = Pt(11)

        # Recommendation image or text
        if device.recommendation_crop_path and Path(device.recommendation_crop_path).exists():
            p_img = cell.add_paragraph()
            p_img.add_run().add_picture(str(device.recommendation_crop_path), width=Inches(5.5))
        elif device.recommendations:
            for rec in device.recommendations:
                p_r = cell.add_paragraph()
                r_item = p_r.add_run(f"• {rec.item}: ")
                r_item.bold = True
                r_item.font.name = "Calibri"
                r_rem = p_r.add_run(rec.remarks)
                r_rem.font.name = "Calibri"

    def _populate_checklist_tables(self, tables: List, device: CheckPointDeviceData) -> None:
        # Build lookup by item number
        items_by_no = {item.no: item for item in device.checklist}

        for t in tables:
            for r in t.rows:
                col0 = r.cells[0].text.strip()
                if col0.isdigit():
                    num = int(col0)
                    if num in items_by_no:
                        item = items_by_no[num]
                        # Col 3: Pass? (Y/N)
                        if len(r.cells) >= 4:
                            r.cells[3].text = item.pass_status
                        # Col 4: Remarks
                        if len(r.cells) >= 5:
                            r.cells[4].text = item.remarks

    def _populate_appendix_tables(self, tables: List, device: CheckPointDeviceData) -> None:
        if not tables:
            return

        # Target single or multiple appendix tables
        # Group appendix images by appendix number
        images_by_app = {}
        for img in device.appendix_images:
            if img.appendix_no not in images_by_app:
                images_by_app[img.appendix_no] = []
            images_by_app[img.appendix_no].append(img)

        # Total appendices
        max_app = 13 if device.system.is_gateway else 12

        # Use first table's first cell to populate cleanly
        target_cell = tables[0].rows[0].cells[0]

        # Clear existing elements in all appendix tables
        for t in tables:
            for r in t.rows:
                for c in r.cells:
                    p0 = c.paragraphs[0]
                    p0.text = ""
                    for p in list(c.paragraphs[1:]):
                        p._p.getparent().remove(p._p)

        # Now populate target_cell with each appendix in order
        for app_num in range(1, max_app + 1):
            imgs = images_by_app.get(app_num, [])
            for img in imgs:
                if Path(img.image_path).exists():
                    p_img = target_cell.add_paragraph()
                    try:
                        p_img.add_run().add_picture(str(img.image_path), width=Inches(5.5))
                    except Exception:
                        pass

            # Title
            title = self._get_appendix_title(app_num, device.system.is_gateway)
            p_cap = target_cell.add_paragraph()
            r_cap = p_cap.add_run(f"\n{title}\n")
            r_cap.bold = True
            r_cap.font.name = "Calibri"
            r_cap.font.size = Pt(11)

    def _get_appendix_title(self, app_num: int, is_gateway: bool) -> str:
        titles_gw = {
            1: "Appendix 1 – Version Checking",
            2: "Appendix 2 – Interfaces",
            3: "Appendix 3 – Routing Information",
            4: "Appendix 4 – Power Supply & Hardware Sensors",
            5: "Appendix 5 – DNS Information",
            6: "Appendix 6 – Date and Time",
            7: "Appendix 7 – CPU Utilization",
            8: "Appendix 8 – Disk Usage",
            9: "Appendix 9 – Log Storage",
            10: "Appendix 10 – Network Performance",
            11: "Appendix 11 – Interface Performance Statistics",
            12: "Appendix 12 – License Status",
            13: "Appendix 13 – Backup Configuration",
        }
        titles_mgmt = {
            1: "Appendix 1 – Version Checking",
            2: "Appendix 2 – Interfaces",
            3: "Appendix 3 – Routing Information",
            4: "Appendix 4 – Power Supply & Hardware Sensors",
            5: "Appendix 5 – DNS Information",
            6: "Appendix 6 – Date and Time",
            7: "Appendix 7 – CPU Utilization",
            8: "Appendix 8 – Disk Usage",
            9: "Appendix 9 – Log Storage",
            10: "Appendix 10 – Interface Performance Statistics",
            11: "Appendix 11 – License Status",
            12: "Appendix 12 – Backup Configuration",
        }
        return titles_gw.get(app_num, f"Appendix {app_num}") if is_gateway else titles_mgmt.get(app_num, f"Appendix {app_num}")
