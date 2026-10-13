"""
Report Generator for Trend Micro Deep Security PM Report.
Populates official Word (.docx) report using non-destructive in-place editing,
preserving corporate headers, footers, styles, and FORMCHECKBOX states.
"""

import shutil
from pathlib import Path
from typing import Optional, Dict, List
import docx
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from .models import DeepSecurityReportData, DeepSecurityChecklistItem
from ..report.images import replace_appendix_image

W_TXBX = qn("w:txbxContent")
W_P = qn("w:p")
W_T = qn("w:t")

# Map of Appendix Number -> primary embedded rId in template
APPENDIX_RID_MAP: Dict[int, List[str]] = {
    0: ["rId25"],            # Item 1 Version overview
    1: ["rId26", "rId27", "rId28"],  # Dashboard, Widgets, Add Widgets
    2: ["rId29"],            # Alerts
    3: ["rId30", "rId31"],  # Events & Reports, CSV Export
    4: ["rId32"],            # License Status
    5: ["rId33"],            # Agent Inventory
    6: ["rId34", "rId35"],  # Updates Security Summary, Components Update
    7: ["rId36"],            # System Information
    8: ["rId37", "rId38", "rId39"],  # Software update / Server health
    9: ["rId40"],            # Health check / DSM Version
}

class DeepSecurityCTCReportGenerator:
    """
    Non-destructive DOCX generator for Trend Micro Deep Security PM reports.
    """

    def __init__(self, default_template: Optional[Path] = None):
        if default_template is None:
            # Navigate from backend/app/services/deep_security/generator.py up to repo root
            repo_root = Path(__file__).resolve().parent.parent.parent.parent.parent
            self.default_template = repo_root / "templates" / "deep_security_pm_template.docx"
            if not self.default_template.exists():
                # Fallback to backend/templates if present
                self.default_template = repo_root / "backend" / "templates" / "deep_security_pm_template.docx"
        else:
            self.default_template = Path(default_template)

    def generate_report(
        self,
        report_data: DeepSecurityReportData,
        output_path: Path,
        template_path: Optional[Path] = None,
    ) -> Path:
        tmpl = template_path or self.default_template
        if not tmpl.exists():
            raise FileNotFoundError(f"Template not found at: {tmpl}")

        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(str(tmpl), str(output_path))

        doc = Document(str(output_path))

        # 1. Update cover page textboxes
        self._update_cover_textboxes(doc, report_data)

        # 2. Update author details (Table 0)
        self._update_author_details(doc, report_data)

        # 3. Update appliance info (Table 2)
        self._update_appliance_table(doc, report_data)

        # 4. Update executive summary (Table 3)
        self._update_executive_summary(doc, report_data)

        # 5. Update version item (Table 4)
        self._update_version_table(doc, report_data)

        # 6. Update checklist tables (Tables 5, 6, 7)
        self._update_checklist_tables(doc, report_data)

        # 7. Update appendix images
        self._update_appendix_images(doc, report_data)

        doc.save(str(output_path))
        return output_path

    def _update_cover_textboxes(self, doc: Document, report_data: DeepSecurityReportData) -> None:
        for p in doc.paragraphs[:30]:
            for tb in p._p.iter(W_TXBX):
                for tb_p in tb.iter(W_P):
                    for t_el in tb_p.iter(W_T):
                        txt = t_el.text or ""
                        # Replace Customer Name
                        if "LQ Retail Sdn Bhd" in txt or "The Exchange TRX" in txt:
                            t_el.text = report_data.customer
                        # Replace PM Date
                        elif "PM DATE" in txt and ":" in txt:
                            t_el.text = f"PM DATE       : {report_data.pm_date}"
                        # Replace FSMS ID
                        elif "FSMS ID" in txt and ":" in txt:
                            t_el.text = f"FSMS ID       : {report_data.fsms_id}"
                        # Replace Contract No
                        elif "Contract No" in txt and ":" in txt:
                            t_el.text = f"Contract No       : {report_data.contract_no}"
                        # Replace Contract Period
                        elif "Contract Period" in txt and ":" in txt:
                            t_el.text = f"Contract Period      : {report_data.contract_period}"

    def _update_author_details(self, doc: Document, report_data: DeepSecurityReportData) -> None:
        if len(doc.tables) < 1:
            return
        t0 = doc.tables[0]
        for r_idx, r in enumerate(t0.rows):
            if len(r.cells) >= 2:
                lbl = r.cells[0].text.strip().lower()
                # Author
                if lbl == "name" and r_idx < 5:
                    r.cells[1].text = report_data.author_name
                elif lbl == "designation" and r_idx < 5:
                    r.cells[1].text = report_data.author_designation
                elif lbl == "date" and r_idx < 5:
                    r.cells[1].text = report_data.author_date
                elif lbl == "signature" and r_idx < 5 and report_data.author_signature_text:
                    r.cells[1].text = report_data.author_signature_text
                # Verified by
                elif lbl == "name" and 5 <= r_idx < 10:
                    r.cells[1].text = report_data.verified_name
                elif lbl == "designation" and 5 <= r_idx < 10:
                    r.cells[1].text = report_data.verified_designation
                elif lbl == "date" and 5 <= r_idx < 10:
                    r.cells[1].text = report_data.verified_date

    def _update_appliance_table(self, doc: Document, report_data: DeepSecurityReportData) -> None:
        if len(doc.tables) < 3:
            return
        t2 = doc.tables[2]
        for r in t2.rows:
            if len(r.cells) >= 2:
                lbl = r.cells[0].text.strip().lower()
                if "model" in lbl:
                    r.cells[1].text = report_data.system.model
                elif "hostname" in lbl:
                    r.cells[1].text = report_data.system.hostname

    def _update_executive_summary(self, doc: Document, report_data: DeepSecurityReportData) -> None:
        if len(doc.tables) < 4:
            return
        t3 = doc.tables[3]
        cell = t3.rows[0].cells[0]

        # In-place paragraph update to retain XML cell properties
        paras = cell.paragraphs
        if len(paras) >= 5 and report_data.key_highlights:
            # P1..P4 correspond to highlights
            for i, hl in enumerate(report_data.key_highlights[:4]):
                p_idx = i + 1
                if p_idx < len(paras):
                    paras[p_idx].text = hl

        # Update recommendations
        recs = report_data.recommendations
        if recs:
            # Manager version recommendation (P7)
            if len(paras) >= 8 and len(recs) >= 1:
                paras[7].text = recs[0].remarks.split("\n")[0]
            # Agent version recommendation (P13)
            if len(paras) >= 14 and len(recs) >= 2:
                paras[13].text = recs[1].remarks

    def _update_version_table(self, doc: Document, report_data: DeepSecurityReportData) -> None:
        if len(doc.tables) < 5:
            return
        t4 = doc.tables[4]
        for r in t4.rows:
            if len(r.cells) >= 3 and r.cells[0].text.strip() == "1":
                r.cells[2].text = f"Version/build: {report_data.system.manager_version}"

    def _update_checklist_tables(self, doc: Document, report_data: DeepSecurityReportData) -> None:
        # Tables 5, 6, 7 are checklist tables
        checklist_tables = [doc.tables[idx] for idx in [5, 6, 7] if idx < len(doc.tables)]
        
        # Build lookup by task text
        subtask_map = {}
        for item in report_data.checklist:
            for st in item.subtasks:
                subtask_map[st.task.strip().lower()] = st

        for t in checklist_tables:
            for r in t.rows:
                if len(r.cells) >= 6:
                    task_col = r.cells[1].text.strip().lower()
                    matched_subtask = None
                    for k, st in subtask_map.items():
                        if k in task_col or task_col in k:
                            matched_subtask = st
                            break

                    if matched_subtask:
                        # 1. Update Result FORMCHECKBOX (Cell 4)
                        self._set_result_checkbox(r.cells[4], matched_subtask.result)
                        # 2. Update Remarks (Cell 5)
                        if matched_subtask.remarks:
                            r.cells[5].text = matched_subtask.remarks

    def _set_result_checkbox(self, cell, result_value: str) -> None:
        """Sets FORMCHECKBOX state for Yes/No in a table cell."""
        is_yes = (result_value.strip().lower() == "yes")
        cbs = cell._tc.xpath(".//w:checkBox")
        if len(cbs) >= 2:
            def0 = cbs[0].find(qn("w:default"))
            def1 = cbs[1].find(qn("w:default"))
            if def0 is not None:
                def0.set(qn("w:val"), "true" if is_yes else "false")
            if def1 is not None:
                def1.set(qn("w:val"), "false" if is_yes else "true")

    def _update_appendix_images(self, doc: Document, report_data: DeepSecurityReportData) -> None:
        """
        Replaces appendix image slots with provided appendix screenshots.
        """
        # Map images by appendix number
        images_by_app: Dict[int, List[Path]] = {}
        for img in report_data.appendix_images:
            p = Path(img.image_path)
            if p.exists():
                images_by_app.setdefault(img.appendix_no, []).append(p)

        for app_num, paths in images_by_app.items():
            rids = APPENDIX_RID_MAP.get(app_num, [])
            for idx, img_path in enumerate(paths):
                if idx < len(rids):
                    replace_appendix_image(doc, rids[idx], img_path)
