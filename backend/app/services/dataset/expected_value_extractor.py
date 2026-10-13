"""
Expected value extractor for PM Report Dataset Importer.
Extracts ground truth PMData fields and read-only checkbox states from completed DOCX reports.
"""

import re
from typing import Dict, Any, List, Optional, Tuple
import docx
from docx.table import _Cell

W14_VAL = "{http://schemas.microsoft.com/office/word/2010/wordml}val"
W_TXBX = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}txbxContent"
W_P = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p"
W_T = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"

CHECKED_GLYPH = "\u2612"    # ☒

def normalize_status(val: Optional[str]) -> Optional[str]:
    if not val:
        return None
    v = val.strip().lower()
    v = re.sub(r'[\s\/\-]+', '_', v)
    if v in ("n_a", "na"):
        return "n_a"
    if v in ("un_available", "unavailable"):
        return "unavailable"
    return v

def get_cell_checkbox_status(cell: _Cell, valid_options: List[str]) -> Optional[str]:
    """
    Reads the selected option from SDT checkbox controls within a cell.
    Returns normalized option name or None.
    """
    for p in cell.paragraphs:
        # Extract full text including any runs inside SDT elements
        p_text_lower = "".join(el.text or "" for el in p._p.iter(W_T)).strip().lower()
        if not p_text_lower:
            p_text_lower = p.text.strip().lower()
        for sdt in p._p.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}sdt"):
            is_checked = False
            for el in sdt.iter():
                tag = el.tag.split("}")[-1]
                # Check w14:val attribute or any val attribute
                val_attr = el.attrib.get(W14_VAL) or el.attrib.get("val") or el.attrib.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val")
                if tag == "checked" and val_attr in ("1", "true"):
                    is_checked = True
                elif tag == "t" and el.text and CHECKED_GLYPH in el.text:
                    is_checked = True

            if is_checked:
                for opt in valid_options:
                    if opt.lower() in p_text_lower:
                        return normalize_status(opt)

    return None

def extract_expected_values(doc: docx.Document) -> Tuple[Dict[str, Any], List[str]]:
    """
    Extracts expected values and statuses from the completed DOCX report.
    Returns (expected_dict, warnings).
    """
    warnings: List[str] = []

    expected: Dict[str, Any] = {
        "metadata": {
            "pm_date": None,
            "fsms_id": None,
            "contract_number": None,
            "contract_period": None,
            "customer": None,
        },
        "application": {
            "model": None,
            "hostname": None,
            "serial_primary": None,
            "serial_secondary": None,
        },
        "hardware": {
            "hard_disk_status": None,
            "cpu_usage": None,
            "cpu_status": None,
            "memory_usage": None,
            "memory_status": None,
            "active_ports": None,
            "interface_status": None,
        },
        "licensing": {
            "forticare_status": None,
            "expiration_date": None,
            "antivirus_status": None,
            "ips_status": None,
            "web_filter_status": None,
            "anti_spam_status": None,
            "other_status": None,
        },
        "firmware": {
            "current": None,
            "build": None,
            "maturity": None,
            "recommended": None,
            "status": None,
            "ffr_status": None,
            "eos_status": None,
        },
        "vpn": {
            "total": None,
            "active": None,
            "inactive": None,
        },
        "configuration": {
            "uptime": None,
            "gui_access": None,
            "cli_access": None,
            "signature_status": None,
            "timezone": None,
            "datetime_status": None,
            "internet_access": None,
            "ha_status": None,
            "event_logging": None,
            "log_manager": None,
            "backup_status": None,
        }
    }

    def _extract_meta(clean: str):
        if "pm date:" in clean.lower() and not expected["metadata"]["pm_date"]:
            m = re.search(r"PM DATE:\s*([0-9\s\/]+)", clean, re.IGNORECASE)
            if m:
                expected["metadata"]["pm_date"] = re.sub(r"\s+", "", m.group(1))

        if ("fsms" in clean.lower() and "id:" in clean.lower()) and not expected["metadata"]["fsms_id"]:
            m = re.search(r"FSMS\s*ID:\s*([0-9]+)", clean, re.IGNORECASE)
            if m:
                expected["metadata"]["fsms_id"] = m.group(1).strip()

        if "contract no" in clean.lower() and not expected["metadata"]["contract_number"]:
            m = re.search(r"CONTRACT NO\s*:\s*(\S+)", clean, re.IGNORECASE)
            if m:
                expected["metadata"]["contract_number"] = m.group(1).strip()

        if "contract period:" in clean.lower() and not expected["metadata"]["contract_period"]:
            m = re.search(r"CONTRACT PERIOD:\s*(.+)", clean, re.IGNORECASE)
            if m:
                expected["metadata"]["contract_period"] = re.sub(r"(\d)\s+(\d)", r"\1\2", m.group(1).strip())

        if "aeon credit" in clean.lower() and not expected["metadata"]["customer"]:
            expected["metadata"]["customer"] = "AEON Credit Services (M) Berhad"
        elif ("customer:" in clean.lower() or "client:" in clean.lower()) and not expected["metadata"]["customer"]:
            m = re.search(r"(?:customer|client):\s*(.+)", clean, re.IGNORECASE)
            if m:
                expected["metadata"]["customer"] = m.group(1).strip()

    # 1. Extract cover page metadata from textboxes & paragraphs
    for p in doc.paragraphs[:50]:
        # Check textboxes inside paragraph
        for tb in p._p.iter(W_TXBX):
            for tb_p in tb.iter(W_P):
                t_elems = [el.text or "" for el in tb_p.iter(W_T)]
                txt = "".join(t_elems).strip()
                _extract_meta(re.sub(r"\s+", " ", txt))

        # Check paragraph text itself
        p_txt = p.text.strip()
        if p_txt:
            _extract_meta(re.sub(r"\s+", " ", p_txt))

    # Also check Table 0 (Author's Details) for date if missing
    if not expected["metadata"]["pm_date"] and len(doc.tables) > 0:
        t0 = doc.tables[0]
        for r in t0.rows:
            if len(r.cells) >= 3 and "date" in r.cells[0].text.lower():
                d_val = r.cells[2].text.strip()
                if d_val:
                    expected["metadata"]["pm_date"] = d_val
                    break

    # 2. Extract application and PM body tables
    for t_idx, t in enumerate(doc.tables):
        for r_idx, r in enumerate(t.rows):
            cell_texts = [c.text.strip() for c in r.cells]
            if not cell_texts:
                continue

            c0 = cell_texts[0].lower()
            c1 = cell_texts[1] if len(cell_texts) > 1 else ""

            # Application Table (Model, Hostname, Serial Number)
            if c0 == "model" and not expected["application"]["model"]:
                expected["application"]["model"] = c1
            elif c0 == "hostname" and not expected["application"]["hostname"]:
                expected["application"]["hostname"] = c1
            elif "serial" in c0 and not expected["application"]["serial_primary"]:
                m_prim = re.search(r"([A-Z0-9]+)\s*\(Primary\)", c1)
                if m_prim:
                    expected["application"]["serial_primary"] = m_prim.group(1)
                m_sec = re.search(r"([A-Z0-9]+)\s*\(Secondary\)", c1)
                if m_sec:
                    expected["application"]["serial_secondary"] = m_sec.group(1)
                if not expected["application"]["serial_primary"] and c1:
                    expected["application"]["serial_primary"] = c1.split(",")[0].strip()

            # Health checklist tables with >= 4 cells (Item, Task, Description, Status, Remarks)
            if len(r.cells) >= 4:
                item_code = cell_texts[0].strip()
                task_name = cell_texts[1].strip()
                status_cell = r.cells[3]
                remarks_cell = r.cells[4] if len(r.cells) >= 5 else None
                remarks_text = remarks_cell.text.strip() if remarks_cell else ""
                row_key = f"{item_code} {task_name}".lower()

                # --- 1. Hardware ---
                if "hard disk" in row_key or item_code == "1.1":
                    expected["hardware"]["hard_disk_status"] = get_cell_checkbox_status(
                        status_cell, ["Healthy", "Unhealthy", "N/A"]
                    )
                elif "cpu usage" in row_key or item_code == "1.2":
                    expected["hardware"]["cpu_status"] = get_cell_checkbox_status(
                        status_cell, ["Healthy", "Unhealthy", "N/A"]
                    )
                    m = re.search(r"cpu usage:\s*(\d+)%", remarks_text, re.IGNORECASE)
                    if m:
                        expected["hardware"]["cpu_usage"] = int(m.group(1))
                elif "memory usage" in row_key or item_code == "1.3":
                    expected["hardware"]["memory_status"] = get_cell_checkbox_status(
                        status_cell, ["Healthy", "Unhealthy", "N/A"]
                    )
                    m = re.search(r"memory usage:\s*(\d+)%", remarks_text, re.IGNORECASE)
                    if m:
                        expected["hardware"]["memory_usage"] = int(m.group(1))
                elif "interface monitoring" in row_key or item_code == "1.4":
                    expected["hardware"]["interface_status"] = get_cell_checkbox_status(
                        status_cell, ["Healthy", "Unhealthy", "N/A"]
                    )
                    m = re.search(r"ports in-used:\s*(\d+)", remarks_text, re.IGNORECASE)
                    if m:
                        expected["hardware"]["active_ports"] = int(m.group(1))

                # --- 2. Licensing ---
                elif "forticare" in row_key or item_code == "2.1":
                    expected["licensing"]["forticare_status"] = get_cell_checkbox_status(
                        status_cell, ["Valid", "Expired", "N/A"]
                    )
                    m = re.search(r"expiration date:\s*([0-9\/\-]+)", remarks_text, re.IGNORECASE)
                    if m:
                        expected["licensing"]["expiration_date"] = m.group(1).strip()
                elif "utm bundle" in row_key or item_code == "2.2":
                    # Check individual UTM features inside remarks cell paragraphs
                    if remarks_cell:
                        for p in remarks_cell.paragraphs:
                            p_txt = p.text.strip().lower()
                            for feat in ["antivirus", "ips", "web filtering", "anti-spam", "others"]:
                                if feat in p_txt:
                                    field_name = feat.replace(" ", "_").replace("-", "_")
                                    if field_name == "web_filtering":
                                        field_name = "web_filter"
                                    elif field_name == "others":
                                        field_name = "other"
                                    key = f"{field_name}_status"

                                    # Check SDTs in this paragraph
                                    sdt_checked_list = []
                                    for sdt in p._p.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}sdt"):
                                        chk = any(
                                            el.attrib.get(W14_VAL) in ("1", "true")
                                            for el in sdt.iter()
                                            if el.tag.split("}")[-1] == "checked"
                                        ) or any(
                                            el.text == CHECKED_GLYPH
                                            for el in sdt.iter()
                                            if el.tag.split("}")[-1] == "t"
                                        )
                                        sdt_checked_list.append(chk)

                                    if len(sdt_checked_list) >= 3:
                                        if sdt_checked_list[0]:
                                            expected["licensing"][key] = "valid"
                                        elif sdt_checked_list[1]:
                                            expected["licensing"][key] = "expired"
                                        elif sdt_checked_list[2] or (len(sdt_checked_list) > 3 and sdt_checked_list[3]):
                                            expected["licensing"][key] = "n_a"

                # --- 3. Firmware ---
                elif "firewall firmware version" in row_key or item_code == "3.1":
                    expected["firmware"]["status"] = get_cell_checkbox_status(
                        status_cell, ["Up-to-Date", "Out-of-Date"]
                    )
                    m_curr = re.search(
                        r"current firmware(?: version)?:\s*(?:fortios\s*)?(v?[\d\.]+)\s*(?:build(\d+))?\s*(?:\(([^)]+)\))?",
                        remarks_text,
                        re.IGNORECASE
                    )
                    if m_curr:
                        expected["firmware"]["current"] = m_curr.group(1).lstrip("v")
                        expected["firmware"]["build"] = m_curr.group(2) if m_curr.group(2) else None
                        expected["firmware"]["maturity"] = m_curr.group(3) if m_curr.group(3) else None

                    m_lat = re.search(
                        r"(?:latest(?: firmware(?: version)?)?|newer release)[\s:]*(v?[\d\.]+)",
                        remarks_text,
                        re.IGNORECASE
                    )
                    if m_lat:
                        expected["firmware"]["recommended"] = m_lat.group(1).lstrip("v")

                elif "final firmware release" in row_key or item_code == "3.2":
                    expected["firmware"]["ffr_status"] = get_cell_checkbox_status(
                        status_cell, ["Valid", "Expired", "N/A"]
                    )
                elif "end of support" in row_key or item_code == "3.3":
                    expected["firmware"]["eos_status"] = get_cell_checkbox_status(
                        status_cell, ["Valid", "Expired", "N/A"]
                    )

                # --- 4. VPN ---
                elif "check number of tunnel" in row_key or item_code == "4.1":
                    st_text = status_cell.text.strip()
                    m_tot = re.search(r"total vpn tunnel\(s\):\s*(\d+)", st_text, re.IGNORECASE)
                    m_act = re.search(r"active vpn tunnel\(s\):\s*(\d+)", st_text, re.IGNORECASE)
                    m_inact = re.search(r"inactive vpn tunnel\(s\):\s*(\d+)", st_text, re.IGNORECASE)
                    if m_tot:
                        expected["vpn"]["total"] = int(m_tot.group(1))
                    if m_act:
                        expected["vpn"]["active"] = int(m_act.group(1))
                    if m_inact:
                        expected["vpn"]["inactive"] = int(m_inact.group(1))

                # --- 6. Configuration ---
                elif "firewall system uptime" in row_key or item_code == "6.1":
                    st_text = status_cell.text.strip()
                    m = re.search(r"uptime:\s*(.+)", st_text, re.IGNORECASE)
                    expected["configuration"]["uptime"] = m.group(1).strip() if m else (st_text or None)
                elif "accessibility of fortigate via gui" in row_key or item_code == "6.2":
                    expected["configuration"]["gui_access"] = get_cell_checkbox_status(
                        status_cell, ["Pass", "Failed", "N/A"]
                    )
                elif "accessibility of fortigate via cli" in row_key or item_code == "6.3":
                    expected["configuration"]["cli_access"] = get_cell_checkbox_status(
                        status_cell, ["Pass", "Failed", "N/A"]
                    )
                elif "signatures version" in row_key or item_code == "6.4":
                    expected["configuration"]["signature_status"] = get_cell_checkbox_status(
                        status_cell, ["Up-to-Date", "Out-of-Date", "N/A"]
                    )
                elif "firewall date & time" in row_key or item_code == "6.5":
                    expected["configuration"]["datetime_status"] = get_cell_checkbox_status(
                        status_cell, ["Pass", "Failed", "N/A"]
                    )
                    m = re.search(r"time zone is set as\s*(.+)", remarks_text, re.IGNORECASE)
                    if m:
                        expected["configuration"]["timezone"] = m.group(1).split("//")[0].strip()
                    elif "gmt +8" in remarks_text.lower():
                        expected["configuration"]["timezone"] = "GMT +8"
                elif "firewall internet access" in row_key or item_code == "6.6":
                    expected["configuration"]["internet_access"] = get_cell_checkbox_status(
                        status_cell, ["Pass", "Failed", "N/A"]
                    )
                elif "high availability" in row_key or item_code == "6.7":
                    expected["configuration"]["ha_status"] = get_cell_checkbox_status(
                        status_cell, ["Pass", "Failed", "N/A"]
                    )
                elif ("logging" in row_key and "event log" in row_key) or item_code == "6.8":
                    expected["configuration"]["event_logging"] = get_cell_checkbox_status(
                        status_cell, ["Available", "Un-available", "Unavailable"]
                    )
                elif "log manager" in row_key or item_code == "6.9":
                    expected["configuration"]["log_manager"] = get_cell_checkbox_status(
                        status_cell, ["Healthy", "Unhealthy", "N/A"]
                    )
                elif "backup configuration" in row_key or item_code == "6.10":
                    expected["configuration"]["backup_status"] = get_cell_checkbox_status(
                        status_cell, ["Passed", "Failed"]
                    )

    # 3. Add warnings for key fields that could not be extracted
    critical_checks = [
        ("metadata.customer", expected["metadata"]["customer"]),
        ("metadata.pm_date", expected["metadata"]["pm_date"]),
        ("application.model", expected["application"]["model"]),
        ("hardware.cpu_usage", expected["hardware"]["cpu_usage"]),
        ("hardware.memory_usage", expected["hardware"]["memory_usage"]),
        ("licensing.forticare_status", expected["licensing"]["forticare_status"]),
        ("firmware.status", expected["firmware"]["status"]),
        ("firmware.current", expected["firmware"]["current"]),
        ("vpn.total", expected["vpn"]["total"]),
        ("configuration.uptime", expected["configuration"]["uptime"]),
    ]

    for field_path, val in critical_checks:
        if val is None:
            warnings.append(f"Unresolved field: {field_path} could not be extracted from report body")

    return expected, warnings
